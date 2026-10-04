# -*- coding: utf-8 -*-
"""开源共享版排版 v2：从《红楼梦全本一百零八回（开源共享版）.md》生成两部DOCX
v2修订：
  1. 修正章回切分正则（补"零/〇"——第101-108回此前失配混入第100回正文）；
  2. 回目目录改为书签超链接（点击跳转至各回；导出PDF后为内部跳转链接）；
  3. 批语占位符已清理（半角?/�→□，专名复原），□=底本录文缺字符号；
  4. 未对位批语附录每条附前后锚文（系于某句之后/之前），标明正文位置。
"""
import re, os, json
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(ROOT, "..", "红楼梦全本一百零八回（开源共享版）.md")
ANN_COLOR = RGBColor(0x8B, 0x00, 0x00)

BYLINE = "西域吃沙的峰兄　撰次"
AI_NOTE = ("成书说明：本项目由西域吃沙的峰兄发起并主持，思路、取舍与总检皆出个人；文字执行由 "
           "GLM-5.3-Flash 大语言模型完成，全程子代理任务88项，累计消耗tokens约6460万。"
           "详见卷首序言与《项目用量统计》。")

def clean(t):
    return re.sub(r'\s+', '', t.replace('&nbsp;','').replace('&quot;','"'))

def set_font(run, name_cn="宋体", size=10.5, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color

def add_para(doc, text, size=10.5, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             cn_font="宋体", bold=False, color=None, space_after=0, pbb=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align; pf.line_spacing = 1.3; pf.space_after = Pt(space_after)
    if pbb: pf.page_break_before = True
    if indent: pf.first_line_indent = Pt(size*2)
    if text:
        r = p.add_run(text); set_font(r, cn_font, size, bold, color)
    return p

def add_heading_chapter(doc, text, bookmark=None, bid=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(18); pf.space_after = Pt(12); pf.line_spacing = 1.3
    r = p.add_run(text); set_font(r, "黑体", 14, bold=True, color=RGBColor(0,0,0))
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'0'); pPr.append(ol)
    if bookmark:
        bs = OxmlElement('w:bookmarkStart'); bs.set(qn('w:id'), str(bid)); bs.set(qn('w:name'), bookmark)
        be = OxmlElement('w:bookmarkEnd');   be.set(qn('w:id'), str(bid))
        p._p.insert(0, bs); p._p.append(be)
    return p

def add_toc_link(cell_para, text, anchor, size=9):
    hl = OxmlElement('w:hyperlink'); hl.set(qn('w:anchor'), anchor); hl.set(qn('w:history'), '1')
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rF = OxmlElement('w:rFonts'); rF.set(qn('w:eastAsia'), '宋体'); rF.set(qn('w:ascii'), 'Times New Roman')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size*2))
    rPr.append(rF); rPr.append(sz); r.append(rPr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    hl.append(r); cell_para._p.append(hl)

def add_footer_pagenum(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE \\* arabic \\* MERGEFORMAT')
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = "1"; r.append(t); fld.append(r)
    p._p.append(fld)

def parse_md():
    t = open(MD, encoding="utf-8").read()
    i = re.search(r'^# 第[零〇一二三四五六七八九十百]+回', t, re.M).start()
    body = t[i:]
    parts = re.split(r'^(# 第[零〇一二三四五六七八九十百]+回 .+)$', body, flags=re.M)
    chapters = []
    for k in range(1, len(parts), 2):
        title = parts[k].lstrip('# ').strip()
        content = parts[k+1].strip()
        content = re.sub(r'^-{3,}\s*', '', content).strip()
        blocks = []
        for blk in content.split("\n\n"):
            blk = blk.strip()
            if not blk or re.fullmatch(r'-{3,}', blk): continue
            is_verse = "\n" in blk
            if is_verse:
                lines = [l.strip().lstrip('　') for l in blk.split("\n") if l.strip()]
                blocks.append({"verse": True, "text": clean(blk), "lines": lines})
            else:
                blocks.append({"verse": False, "text": clean(blk), "lines": [clean(blk)]})
        chapters.append({"title": title, "blocks": blocks})
    return chapters

def toc_titles(chapters):
    out=[]
    for c in chapters:
        m = re.match(r'第([\d零〇一二三四五六七八九十百]+)回 (.+)', c["title"])
        out.append("第%s回　%s" % (m.group(1), m.group(2)) if m else c["title"])
    return out

def title_page(doc):
    for _ in range(4): add_para(doc, "", indent=False)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("红楼梦"), "黑体", 36, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("一百零八回足本（开源共享版）"), "楷体", 16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    set_font(p.add_run(BYLINE), "楷体", 13)
    add_para(doc, "", indent=False)
    add_para(doc, AI_NOTE, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

def xu_page(doc):
    xu = open(os.path.join(ROOT, "..", "文档集", "06_序言.md"), encoding="utf-8").read().strip()
    lines = xu.split("\n")
    body = [l.strip() for l in lines[1:] if l.strip()]
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.page_break_before = True
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.space_after = Pt(14); pf.line_spacing = 1.3
    set_font(p.add_run("序"), "黑体", 16, bold=True, color=RGBColor(0,0,0))
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'0'); pPr.append(ol)
    bs = OxmlElement('w:bookmarkStart'); bs.set(qn('w:id'), '998'); bs.set(qn('w:name'), 'xu')
    be = OxmlElement('w:bookmarkEnd'); be.set(qn('w:id'), '998')
    p._p.insert(0, bs); p._p.append(be)
    for l in body:
        if l.startswith("**") and l.endswith("**"):
            add_para(doc, l.strip("*"), size=11, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="楷体")
        elif l.startswith("## "):
            add_para(doc, l[3:], size=12, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="黑体", bold=True)
        else:
            add_para(doc, l, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

def toc_block(doc, titles):
    p = add_para(doc, "回目目录", size=14, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                 cn_font="黑体", bold=True, space_after=8, pbb=True)
    p2 = add_para(doc, "（目录各条已设内部链接，点击可跳转至该回；PDF阅读器侧边栏亦含全书书签）", size=8.5,
                  indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="楷体", color=RGBColor(0x60,0x60,0x60))
    table = doc.add_table(rows=(len(titles)+1)//2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    i = 0
    for r in table.rows:
        for c in r.cells:
            if i < len(titles):
                c.text = ""
                add_toc_link(c.paragraphs[0], titles[i], "ch%d" % (i+1))
                i += 1

def build(out_path, annos_mode=False):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6); sec.right_margin = Cm(2.6)
    add_footer_pagenum(doc)
    chapters = parse_md()
    assert len(chapters) == 108, "章回切分数异常: %d" % len(chapters)
    title_page(doc)
    if annos_mode:
        data = json.load(open(os.path.join(ROOT,"data","脂批对位.json"), encoding="utf-8"))
        unm = json.load(open(os.path.join(ROOT,"data","未对位批语.json"), encoding="utf-8"))
        add_para(doc, "本册为前八十回脂批对照本：4089条脂批，对位嵌入3860条（〔〕楷体小字随文），余229条附卷末并注所系句位；批语中“□”为底本缺字符号，全角“？”为批语原有语气。体例详见《02_脂批对照本说明》。",
                 size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    xu_page(doc)
    toc_block(doc, toc_titles(chapters))
    for idx, ch in enumerate(chapters):
        n = idx+1
        add_heading_chapter(doc, ch["title"], bookmark="ch%d" % n, bid=n)
        annos = (data[str(n)]["annos"] if annos_mode and n<=80 and str(n) in data else [])
        ai = 0; cursor = 0
        for blk in ch["blocks"]:
            start, end = cursor, cursor+len(blk["text"]); cursor = end
            in_range = [a for a in annos[ai:] if a["pos"] < end]
            ai += len(in_range)
            if blk["verse"]:
                for ln in blk["lines"]:
                    p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.3
                    set_font(p.add_run(ln), "楷体", 10.5)
                for a in in_range:
                    add_para(doc, "〔批："+a["batch"]+"〕", size=8.5, indent=False,
                             align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="楷体", color=ANN_COLOR)
                continue
            cuts = [max(0, min(a["pos"]-start, len(blk["text"]))) for a in in_range]
            segs=[]; last=0
            for cp in cuts: segs.append(blk["text"][last:cp]); last=cp
            segs.append(blk["text"][last:])
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf.line_spacing = 1.3
            pf.first_line_indent = Pt(21)
            for k, seg in enumerate(segs):
                if seg: set_font(p.add_run(seg), "宋体", 10.5)
                if k < len(in_range):
                    set_font(p.add_run("〔"+in_range[k]["batch"]+"〕"), "楷体", 8.5, color=ANN_COLOR)
    if annos_mode:
        add_heading_chapter(doc, "附录：未对位批语辑录（229条）", bookmark="appendix", bid=999)
        add_para(doc, "以下批语因正文对应段已被鬼本织入文字替换或锚文变动，未能随文排出，按回辑录。每条后附其原所系之前后锚文（据脂本原录文），读者可据此在脂本原书中复核其位置。", indent=False)
        by = defaultdict(list)
        for u in unm: by[u["回"]].append(u)
        for n2 in sorted(by):
            add_para(doc, "第%d回" % n2, size=10.5, indent=False, cn_font="黑体", bold=True)
            for u in by[n2]:
                add_para(doc, "〔"+u["batch"]+"〕", size=8.5, indent=False, cn_font="楷体", color=ANN_COLOR)
                hint = "（系于「" + u["before"] + "」句后" if u["before"] else "（"
                if u["after"]: hint += "、「" + u["after"] + "」句前）" if u["before"] else "系于「" + u["after"] + "」句前）"
                else: hint += "）" if u["before"] else "（回前批）"
                add_para(doc, hint, size=8.5, indent=False, cn_font="宋体", color=RGBColor(0x60,0x60,0x60))
    doc.save(out_path)
    print("saved:", out_path, os.path.getsize(out_path)//1024, "KB")

if __name__ == "__main__":
    build(os.path.join(ROOT, "..", "红楼梦全本一百零八回（开源共享版）.docx"), annos_mode=False)
    build(os.path.join(ROOT, "..", "红楼梦前八十回（脂批对照本·开源共享版）.docx"), annos_mode=True)
